"""Build deterministic synthetic PDF and XLSX files for AI-query QA."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1] / "qa" / "ai-query" / "assets"


def build_pdf() -> bytes:
    buffer = BytesIO()
    page = canvas.Canvas(buffer, pagesize=(612, 792), invariant=1)
    page.setFont("Helvetica", 11)
    lines = [
        "VaultSpace AI QA synthetic technology overview. All entities are fictional.",
        "SolarGrid serves electric utilities and commercial solar operators.",
        "The platform combines load forecasting and battery scheduling.",
        "A fictional integration partner named Marshlight Systems supplies weather feeds.",
        "Battery dispatch optimization is a topic in the planning workflow.",
    ]
    for index, line in enumerate(lines):
        page.drawString(45, 740 - 24 * index, line)
    page.showPage()
    page.save()
    return buffer.getvalue()


def build_xlsx() -> bytes:
    workbook = Workbook()
    fixed_time = datetime(2026, 10, 17, 0, 0, 0, tzinfo=timezone.utc)
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    revenue = workbook.active
    revenue.title = "Revenue"
    revenue.append(["Metric", "2025", "2026"])
    revenue.append(["Revenue USD", 1200000, 1800000])
    revenue.append(["Cost of goods sold USD", 540000, 720000])
    revenue.append(["Gross profit USD", "=B2-B3", "=C2-C3"])
    revenue.append(["Gross margin", "=B4/B2", "=C4/C2"])
    financing = workbook.create_sheet("Financing")
    financing.append(["Metric", "Proposed"])
    financing.append(["New shares", 4000000])
    financing.append(["Price per share USD", 3])
    financing.append(["Gross proceeds USD", "=B2*B3"])
    notes = workbook.create_sheet("Notes")
    notes.append(["All names and figures are fictional QA data."])
    workbook.calculation.fullCalcOnLoad = True
    source = BytesIO()
    workbook.save(source)
    source.seek(0)
    target = BytesIO()
    with ZipFile(source, "r") as src, ZipFile(target, "w", ZIP_DEFLATED, compresslevel=9) as dst:
        for name in sorted(src.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            content = src.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                    b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" '
                    b'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
                    b'xsi:type="dcterms:W3CDTF">2026-10-17T00:00:00Z</dcterms:modified>',
                    content,
                )
            dst.writestr(info, content, compress_type=ZIP_DEFLATED, compresslevel=9)
    return target.getvalue()


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    pdf = ROOT / "technology-overview.pdf"
    xlsx = ROOT / "finance-model.xlsx"
    pdf.write_bytes(build_pdf())
    xlsx.write_bytes(build_xlsx())
    print(f"built {pdf.relative_to(ROOT.parent.parent.parent)} and {xlsx.relative_to(ROOT.parent.parent.parent)}")


if __name__ == "__main__":
    main()
